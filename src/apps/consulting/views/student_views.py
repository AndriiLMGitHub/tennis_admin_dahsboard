# ==========================================
# STUDENT VIEWS
# ==========================================
import io
from datetime import datetime
import logging
from itertools import chain

from django.conf import settings
from django.contrib import messages
from django.core.exceptions import ValidationError, ObjectDoesNotExist
from django.core.files.storage import default_storage
from django.core.paginator import PageNotAnInteger, EmptyPage, Paginator
from django.db import transaction
from django.db.models import Q
from django.http import JsonResponse, HttpResponse, HttpRequest
from django.shortcuts import render, get_object_or_404, redirect
from django.template.defaultfilters import filesizeformat
from django.utils import timezone
from django.utils.translation import gettext_lazy as _
from django.views.decorators.http import require_POST

from apps.consulting.decorators import student_required
from apps.consulting.enums import Specialty
from apps.consulting.forms import CreateRequestForm, NutritionIntakeForm, PsychologyIntakeForm
from apps.consulting.models import Service, Request, Media, Report, PendingUpload, NutritionIntake
from django.contrib.auth import get_user_model

from apps.consulting.storage import generate_r2_upload_service
from apps.consulting.utils import prepare_request_email_payload
from infrastructure.email.tasks import send_email_task

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

User = get_user_model()

TEMPLATE_DIR = 'dashboard/student/requests'

logger = logging.getLogger(__name__)


@student_required
def list_students_requests_view(request: HttpRequest) -> HttpResponse:
    """
    Displays a list of applications for the current student with pagination and filtering.

    The controller receives all applications created by the authorized user,
    filters them by specialty (if provided and valid), optimizes SQL queries
    using select_related, and splits the result into pages.
    """
    # 1. Отримуємо та нормалізуємо параметр фільтрації
    current_specialty = request.GET.get('specialty', '').upper()

    # Визначаємо валідні варіанти. В ідеалі використовувати Specialty.values,
    # якщо використовуєш TextChoices, або список: ['TENNIS', 'NUTRITION', 'PSYCHOLOGY']
    valid_specialties = ['TENNIS', 'NUTRITION', 'PSYCHOLOGY']

    # 2. Формуємо базовий QuerySet (Lazy evaluation - запит у БД ще не йде)
    requests_student_list = (
        Request.objects.filter(user=request.user)
        .select_related('service', 'assigned_specialist')
        .order_by('-created_at')
    )

    # 3. Застосовуємо фільтр ТІЛЬКИ якщо значення легітимне
    if current_specialty in valid_specialties:
        requests_student_list = requests_student_list.filter(specialty_needed=current_specialty)
    else:
        # Якщо в URL передали сміття або нічого, скидаємо змінну,
        # щоб у шаблоні правильно підсвітився таб "All Requests"
        current_specialty = None

    # 4. Пагінація
    student_paginator = Paginator(requests_student_list, settings.ITEMS_PER_PAGE)
    page_number = request.GET.get('page', 1)

    try:
        page_obj = student_paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = student_paginator.page(1)
    except EmptyPage:
        page_obj = student_paginator.page(student_paginator.num_pages)

    breadcrumbs = [
        {'name': _('My Requests'), 'url': None},
    ]

    context = {
        'requests': page_obj,
        'breadcrumbs': breadcrumbs,
        'current_specialty': current_specialty,  # 🌟 Ключовий параметр для фронтенду
    }

    return render(request, f'{TEMPLATE_DIR}/list.html', context)


@student_required
@require_POST
def generate_r2_upload_url(request):
    filename = request.POST.get('filename')
    upload_type = request.POST.get('upload_type', 'general')
    content_type = request.POST.get('content_type', 'application/octet-stream')

    # Безпечно конвертуємо розмір файлу
    try:
        file_size = int(request.POST.get('file_size', 0))
    except (ValueError, TypeError):
        return JsonResponse({'error': _('Invalid file size format.')}, status=400)

    try:
        data = generate_r2_upload_service(
            filename=filename,
            content_type=content_type,
            upload_type=upload_type,
            file_size=file_size
        )
        return JsonResponse(data)

    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=400)

    except Exception as e:
        logger.error(
            f"Failed to generate R2 upload URL for file {filename}. Error: {str(e)}",
            exc_info=True
        )
        return JsonResponse({'error': _('Internal server error occurred.')}, status=500)


@student_required
def create_request_view(request):
    selected_specialty = request.POST.get("specialty_needed", Specialty.TENNIS)
    valid_specialties = {Specialty.TENNIS, Specialty.NUTRITION, Specialty.PSYCHOLOGY}

    base_services = Service.objects.filter(is_active=True)
    service_specialty = selected_specialty if selected_specialty in valid_specialties else Specialty.TENNIS

    form = CreateRequestForm(request.POST or None)
    form.fields["service"].queryset = base_services.filter(specialty=service_specialty)

    nutrition_form_kwargs = {"prefix": "nutr"}
    psychology_form_kwargs = {"prefix": "psych"}

    if request.method == "POST" and selected_specialty == Specialty.NUTRITION:
        nutrition_form_kwargs.update({"data": request.POST, "files": request.FILES})
    elif request.method == "POST" and selected_specialty == Specialty.PSYCHOLOGY:
        psychology_form_kwargs["data"] = request.POST

    nutrition_form = NutritionIntakeForm(**nutrition_form_kwargs)
    psychology_form = PsychologyIntakeForm(**psychology_form_kwargs)

    def create_nutrition_intake(request_obj):
        intake_obj = nutrition_form.save(commit=False)
        intake_obj.student = request.user
        intake_obj.request = request_obj

        r2_file_key = request.POST.get("r2_file_key")
        r2_file_size = request.POST.get('r2_file_size')
        if r2_file_key:
            intake_obj.attachments = r2_file_key
            if r2_file_size:
                intake_obj.file_size_bytes = int(r2_file_size)
            PendingUpload.objects.filter(file_key=r2_file_key).delete()

        intake_obj.save()

    def create_psychology_intake(request_obj):
        psychology_obj = psychology_form.save(commit=False)
        psychology_obj.student = request.user
        psychology_obj.request = request_obj
        psychology_obj.save()

    def create_tennis_media(request_obj):
        media_type = request.POST.get("media_type")
        if media_type not in ("video", "file", "link"):
            return

        media = Media(request=request_obj, media_type=media_type)
        if media_type in ("video", "file"):
            file_key = request.POST.get("uploaded_file_key")
            file_size = request.POST.get("uploaded_file_size")

            if not file_key:
                raise ValidationError(_("Media file was not uploaded successfully."))

            media.file = file_key

            if file_size and file_size.isdigit():
                media.file_size_bytes = int(file_size)

            media.full_clean()
            media.save()
            PendingUpload.objects.filter(file_key=file_key).delete()
            return

        link = form.cleaned_data.get("media_link")
        if not link:
            raise ValidationError(_("Media link is required"))

        media.url = link
        media.full_clean()
        media.save()

    specialty_handlers = {
        Specialty.TENNIS: create_tennis_media,
        Specialty.NUTRITION: create_nutrition_intake,
        Specialty.PSYCHOLOGY: create_psychology_intake,
    }

    if request.method == "POST":
        is_valid = form.is_valid()
        if selected_specialty == Specialty.NUTRITION:
            is_valid = is_valid and nutrition_form.is_valid()
        elif selected_specialty == Specialty.PSYCHOLOGY:
            is_valid = is_valid and psychology_form.is_valid()

        if is_valid:
            try:
                with transaction.atomic():
                    service = form.cleaned_data["service"]

                    if service.specialty != selected_specialty:
                        raise ValidationError(_("Selected service does not match the chosen specialty."))

                    request_obj = Request.objects.create(
                        user=request.user,
                        service=service,
                        specialty_needed=selected_specialty,
                        comment=form.cleaned_data.get("comment", "") if selected_specialty == Specialty.TENNIS else "",
                    )

                    handler = specialty_handlers.get(selected_specialty)
                    if handler:
                        handler(request_obj)

                    admin_emails = list(User.objects.filter(is_superuser=True).values_list("email", flat=True))
                    if admin_emails:
                        current_domain = request.build_absolute_uri("/").rstrip("/")
                        payload = prepare_request_email_payload(
                            request_obj,
                            admin_emails,
                            base_url=current_domain,
                            recipient_role='admin'
                        )
                        transaction.on_commit(lambda: send_email_task.delay(payload))

                messages.success(request, _("Request created successfully"))
                return redirect("student_list_requests")

            except ValidationError as e:
                error_message = e.messages[0] if hasattr(e, 'messages') else str(e)
                messages.error(request, _(f"Validation failed: {error_message}"))

    context = {
        "form": form,
        "nutrition_form": nutrition_form,
        "psychology_form": psychology_form,
        "tennis_services": base_services.filter(specialty=Specialty.TENNIS),
        "nutrition_services": base_services.filter(specialty=Specialty.NUTRITION),
        "psychology_services": base_services.filter(specialty=Specialty.PSYCHOLOGY),
        "services": base_services,
    }

    return render(
        request,
        f"{TEMPLATE_DIR}/create.html",
        context
    )


@student_required
def student_request_detail_view(request, request_id):
    # 1. ОПТИМІЗАЦІЯ: Додаємо 'report' у select_related
    request_item = get_object_or_404(
        Request.objects
        .prefetch_related('media')
        .select_related('service', 'assigned_specialist', 'report'),
        id=request_id,
        user=request.user
    )

    # noinspection PyTypeChecker
    try:
        report = request_item.report
    except ObjectDoesNotExist:
        report = None

    context = {
        'request_item': request_item,
        'report': report,
    }

    return render(request, f'{TEMPLATE_DIR}/detail.html', context)


@student_required
def list_coaches_view(request):
    coaches = User.objects.filter(role="Coach", coach_status="approved").all()

    breadcrumbs = [
        {'name': _('List coaches'), 'url': None},
    ]

    context = {
        'coaches': coaches,
        'breadcrumbs': breadcrumbs,
    }
    return render(request, f'{TEMPLATE_DIR}/list_coaches.html', context)


# ==========================================
# 1. СТОРІНКА СПИСКУ ФАЙЛІВ (GET)
# ==========================================
@student_required
def files_attach_view(request):
    """
    Optimized File Manager View.
    Prevents N+1 storage queries and memory bloat by using .values()
    and resolving URLs only for the current page items.
    """
    # 1. 🚀 ДОСТАЄМО ТІЛЬКИ ТЕКСТ. Жодних об'єктів файлів чи ORM інстансів
    tennis_qs = Media.objects.filter(
        request__user=request.user
    ).values('id', 'media_type', 'file', 'url', 'created_at', 'file_size_bytes')

    nutrition_qs = NutritionIntake.objects.filter(
        student=request.user
    ).exclude(
        Q(attachments__exact='') | Q(attachments__isnull=True)
    ).values('id', 'attachments', 'created_at', 'file_size_bytes',)

    # 2. Уніфікуємо дані в прості Python-словники
    unified_files = []

    for t in tennis_qs:
        is_link = (t['media_type'] == 'link')
        raw_path = t['url'] if is_link else t['file']
        unified_files.append({
            'id': f"t_{t['id']}",
            'size': t['file_size_bytes'],
            'media_type': t['media_type'],  # 'video', 'file', 'link'
            'raw_path': raw_path,
            'date': t['created_at'],
        })

    for n in nutrition_qs:
        unified_files.append({
            'id': f"n_{n['id']}",
            'size': n['file_size_bytes'],
            'media_type': 'document',
            'raw_path': n['attachments'],
            'date': n['created_at'],
        })

    # 3. Сортуємо в оперативній пам'яті (для словників це відбувається за мілісекунду)
    unified_files.sort(key=lambda x: x['date'] or timezone.now(), reverse=True)
    total_count = len(unified_files)

    # 4. РОБИМО ПАГІНАЦІЮ ДО ТОГО, ЯК ЗВЕРНЕТЬСЯ ДО ХМАРИ
    paginator = Paginator(unified_files, 10)  # 10 елементів на сторінку
    page_number = request.GET.get('page', 1)

    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    # 5. 🛡️ DECORATOR: Генеруємо URL ТІЛЬКИ для поточних 10 записів!
    for item in page_obj.object_list:
        if item['media_type'] == 'link':
            item['display_name'] = item['raw_path']
            item['file_url'] = item['raw_path']
            item['size_formatted'] = '—'
        else:
            path = item['raw_path']
            # Дістаємо чисте ім'я файлу (все що після останнього слешу)
            item['display_name'] = path.split('/')[-1] if path else 'Unknown'

            if path:
                try:
                    item['file_url'] = default_storage.url(path)
                except Exception:
                    item['file_url'] = '#'
            else:
                item['file_url'] = '#'

            # Блокуємо HTTP HEAD запит за розміром. Віддаємо заглушку.
            item['size_formatted'] = 'Cloud Storage'

    context = {
        'all_media': page_obj,
        'total_count': total_count,
        'total_bytes': sum(f['size'] for f in unified_files if f.get('size'))
    }
    return render(request, f'{TEMPLATE_DIR}/files/files_attach.html', context)


@student_required
def generate_report_docx_view(request, report_id):
    # Отримуємо звіт
    report = get_object_or_404(Report, id=report_id)
    request_item = report.request
    student = request_item.user

    # Створюємо порожній документ у пам'яті
    doc = Document()

    # --- ЗАГОЛОВОК ---
    title = doc.add_heading('TENNIS VIDEO ANALYSIS', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("Improve Faster with Focused Feedback", style='Subtitle')

    # --- ІНФОРМАЦІЯ ПРО ГРАВЦЯ ---
    doc.add_heading('PLAYER INFORMATION', level=1)
    doc.add_paragraph(f"Full Name: {student.get_full_name() or student.email}")
    doc.add_paragraph(f"Email: {student.email}")

    # --- ЗВІТ КОУЧА ---
    doc.add_heading('COACH ANALYSIS REPORT', level=1)

    if report.overall_assessment:
        doc.add_heading('Overall Assessment', level=2)
        doc.add_paragraph(report.overall_assessment)

    doc.add_heading('Key Strengths', level=2)
    if report.strength_1: doc.add_paragraph(f"1. {report.strength_1}")
    if report.strength_2: doc.add_paragraph(f"2. {report.strength_2}")
    if report.strength_3: doc.add_paragraph(f"3. {report.strength_3}")

    # --- ТОП-3 ПРІОРИТЕТИ ---
    doc.add_heading('TOP 3 PRIORITY IMPROVEMENTS', level=1)

    # Priority 1
    if report.priority_1_issue:
        doc.add_heading('Priority #1', level=2)
        doc.add_paragraph(f"Issue: {report.priority_1_issue}")
        doc.add_paragraph(f"Why it matters: {report.priority_1_why}")
        doc.add_paragraph(f"Recommended drill or exercise: {report.priority_1_drill}")

    # Priority 2
    if report.priority_2_issue:
        doc.add_heading('Priority #2', level=2)
        doc.add_paragraph(f"Issue: {report.priority_2_issue}")
        doc.add_paragraph(f"Why it matters: {report.priority_2_why}")
        doc.add_paragraph(f"Recommended drill or exercise: {report.priority_2_drill}")

    # Priority 3
    if report.priority_3_issue:
        doc.add_heading('Priority #3', level=2)
        doc.add_paragraph(f"Issue: {report.priority_3_issue}")
        doc.add_paragraph(f"Why it matters: {report.priority_3_why}")
        doc.add_paragraph(f"Recommended drill or exercise: {report.priority_3_drill}")

    # --- ПЛАН ДІЙ ---
    if report.action_plan:
        doc.add_heading('2–4 WEEK ACTION PLAN', level=1)
        doc.add_paragraph('Primary practice focus for the next 2–4 weeks:')
        doc.add_paragraph(report.action_plan)

    doc.add_paragraph("\nThank you for choosing our Tennis Video Analysis service.")

    # --- ЗБЕРЕЖЕННЯ В ОПЕРАТИВНУ ПАМ'ЯТЬ ---
    f = io.BytesIO()
    doc.save(f)
    f.seek(0)  # Повертаємо курсор на початок файлу

    # --- ВІДДАЄМО ФАЙЛ  ---
    response = HttpResponse(
        f.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    # Змушуємо браузер завантажити файл із правильним іменем
    response[
        'Content-Disposition'] = f'attachment; filename="Tennis_Analysis_{student.get_short_name() or student.id}.docx"'

    return response
